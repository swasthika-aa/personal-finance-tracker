"""Generate editable PPTX, DOCX, and Markdown from the presentation's JSON source."""

import argparse
import base64
import io
import json
from pathlib import Path

import nbformat
from docx import Document
from docx.shared import Inches as WordInches, Pt as WordPt, RGBColor as WordColor
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


INK = "172C29"
ACCENT = "167464"
MUTED = "536762"
PALE = "EAF3EF"


def extract_figures(notebook_path: Path, destination: Path) -> dict[str, Path]:
    notebook = nbformat.read(notebook_path, as_version=4)
    images = [output["data"]["image/png"] for cell in notebook.cells
              for output in cell.get("outputs", []) if "image/png" in output.get("data", {})]
    names = ["overview", "category_heatmap", "classification", "anomalies", "forecast"]
    if len(images) != len(names):
        raise ValueError("Expected five charts from the executed project notebook; rerun scripts/execute_notebook.py")
    destination.mkdir(parents=True, exist_ok=True)
    figures = {}
    for name, encoded in zip(names, images, strict=True):
        path = destination / f"{name}.png"
        image_bytes = base64.b64decode(encoded)
        with Image.open(io.BytesIO(image_bytes)) as image:
            image.verify()
        path.write_bytes(image_bytes)
        figures[name] = path
    return figures


def markdown_text(content: dict) -> str:
    sections = [f"# {content['title']}\n\n{content['subtitle']}\n"]
    for number, slide in enumerate(content["slides"], start=1):
        parts = [f"## Slide {number}: {slide['title']}\n"]
        parts.extend(f"- {bullet}" for bullet in slide.get("bullets", []))
        if "table" in slide:
            table = slide["table"]
            parts.extend(["", "| " + " | ".join(table["headers"]) + " |",
                          "| " + " | ".join("---" for _ in table["headers"]) + " |"])
            parts.extend("| " + " | ".join(row) + " |" for row in table["rows"])
        if "figure" in slide:
            parts.extend(["", f"![{slide['caption']}](assets/{slide['figure']}.png)", "", slide["caption"]])
        if slide.get("note"):
            parts.extend(["", f"**Note:** {slide['note']}"])
        parts.extend(["", f"**Presenter Notes:** {slide['speaker_notes']}"])
        sections.append("\n".join(parts))
    return "\n\n---\n\n".join(sections) + "\n"


def add_text(slide, text: str, left: float, top: float, width: float, height: float,
             size: int = 20, color: str = INK, bold: bool = False):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.01)
    frame.margin_top = frame.margin_bottom = Inches(0.01)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(color)
    return shape


def add_bullets(slide, bullets: list[str], width: float, size: int, top: float = 1.8):
    shape = slide.shapes.add_textbox(Inches(0.68), Inches(top), Inches(width), Inches(4.6))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0)
    frame.margin_top = frame.margin_bottom = Inches(0)
    for position, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if position == 0 else frame.add_paragraph()
        paragraph.text = f"- {bullet}"
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = RGBColor.from_string(INK)
        paragraph.space_after = Pt(16 if width > 8 else 13)


def add_figure(slide, path: Path):
    with Image.open(path) as image:
        ratio = image.width / image.height
    width = min(7.7, 4.55 * ratio)
    height = width / ratio
    slide.shapes.add_picture(str(path), Inches(4.95 + (7.7 - width) / 2),
                             Inches(1.85 + (4.55 - height) / 2), width=Inches(width), height=Inches(height))


def add_architecture(slide):
    stages = [
        ("01  INPUT", "Bank CSVs\nReceipt JSON"),
        ("02  PREPARE", "Clean records\nQuality audit"),
        ("03  STORE", "SQLite JSON\nPandas mart"),
        ("04  EVALUATE", "Categorization\nAnomalies / forecast"),
        ("05  PRESENT", "Jupyter controls\nLocal reports"),
    ]
    for position, (heading, detail) in enumerate(stages):
        left = 0.68 + position * 2.42
        block = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(2.5), Inches(2.2), Inches(2.0))
        block.fill.solid()
        block.fill.fore_color.rgb = RGBColor.from_string(PALE)
        block.line.fill.background()
        add_text(slide, heading, left + 0.12, 2.75, 1.96, 0.4, 14, ACCENT, True)
        add_text(slide, detail, left + 0.12, 3.3, 1.96, 0.95, 18)
        if position < len(stages) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(left + 2.23), Inches(3.3), Inches(0.16), Inches(0.3))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = RGBColor.from_string(ACCENT)
            arrow.line.fill.background()
    add_text(slide, "One shared pipeline, invoked from Jupyter, the local CLI, or Docker batch.",
             0.68, 5.15, 11.95, 0.8, 22, INK)


def build_powerpoint(content: dict, figures: dict[str, Path], path: Path):
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = content["title"]
    presentation.core_properties.author = "[STUDENT NAME(S)]"
    for number, item in enumerate(content["slides"], start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor.from_string(ACCENT)
        bar.line.fill.background()
        add_text(slide, "PERSONAL FINANCE TRACKER  /  BACHELOR'S PROJECT", 0.68, 0.32, 11.8, 0.3, 11, ACCENT, True)
        add_text(slide, item["title"], 0.68, 0.85, 12, 0.88,
                 28 if len(item["title"]) > 53 else 31, INK, True)
        if item.get("kind") == "cover":
            add_text(slide, content["subtitle"], 0.68, 1.9, 11.8, 0.55, 23, ACCENT)
            add_bullets(slide, item["bullets"], 11.8, 18, top=2.7)
        elif item.get("kind") == "architecture":
            add_architecture(slide)
        elif "figure" in item:
            add_bullets(slide, item.get("bullets", []), 3.95, 18)
            add_figure(slide, figures[item["figure"]])
        elif "table" in item:
            bullets = item.get("bullets", [])
            if bullets:
                add_bullets(slide, bullets, 11.8, 18)
            table_data = item["table"]
            top = 3.05 if bullets else 1.95
            table = slide.shapes.add_table(len(table_data["rows"]) + 1, len(table_data["headers"]),
                                           Inches(0.68), Inches(top), Inches(11.95),
                                           Inches(6.25 - top)).table
            for row_index, row in enumerate([table_data["headers"], *table_data["rows"]]):
                for column_index, text in enumerate(row):
                    cell = table.cell(row_index, column_index)
                    cell.text = text
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(ACCENT if row_index == 0 else PALE if row_index % 2 else "FFFFFF")
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.font.name = "Aptos"
                        paragraph.font.size = Pt(16 if len(table_data["headers"]) == 2 else 14)
                        paragraph.font.bold = row_index == 0
                        paragraph.font.color.rgb = RGBColor.from_string("FFFFFF" if row_index == 0 else INK)
        else:
            add_bullets(slide, item.get("bullets", []), 11.8, 21)
        add_text(slide, item.get("note", ""), 0.68, 6.62, 11.95, 0.5, 12, MUTED)
        add_text(slide, "LOCAL PYTHON  |  JUPYTER  |  REPRODUCIBLE ANALYTICS", 0.68, 7.14, 10.8, 0.2, 9, MUTED)
        add_text(slide, f"{number:02d} / {len(content['slides']):02d}", 11.9, 7.08, 0.8, 0.3, 11, MUTED)
        slide.notes_slide.notes_text_frame.text = item["speaker_notes"]
    presentation.save(path)


def build_word(content: dict, figures: dict[str, Path], path: Path):
    document = Document()
    section = document.sections[0]
    section.page_width = WordInches(8.27)
    section.page_height = WordInches(11.69)
    section.top_margin = section.bottom_margin = WordInches(0.65)
    section.left_margin = section.right_margin = WordInches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = WordPt(11)
    normal.paragraph_format.space_after = WordPt(7)
    for style_name in ["Title", "Heading 1", "Heading 2"]:
        document.styles[style_name].font.color.rgb = WordColor.from_string(ACCENT)
    document.core_properties.title = content["title"]
    document.core_properties.author = "[STUDENT NAME(S)]"
    section.footer.paragraphs[0].text = "Personal Finance Tracker | Bachelor's Project | [STUDENT NAME(S)]"
    for number, item in enumerate(content["slides"], start=1):
        if number > 1:
            document.add_page_break()
        document.add_heading(item["title"], 0 if number == 1 else 1)
        document.add_paragraph(content["subtitle"] if number == 1 else f"Slide {number} of {len(content['slides'])}", "Subtitle")
        for bullet in item.get("bullets", []):
            document.add_paragraph(bullet, style="List Bullet")
        if "table" in item:
            table_data = item["table"]
            table = document.add_table(rows=1, cols=len(table_data["headers"]))
            table.style = "Light Shading Accent 1"
            for cell, heading in zip(table.rows[0].cells, table_data["headers"], strict=True):
                cell.text = heading
            for row in table_data["rows"]:
                for cell, text in zip(table.add_row().cells, row, strict=True):
                    cell.text = text
        if "figure" in item:
            document.add_picture(str(figures[item["figure"]]), width=WordInches(6.6))
            document.add_paragraph(item["caption"], style="Caption")
        if item.get("note"):
            paragraph = document.add_paragraph()
            paragraph.add_run("Note: ").bold = True
            paragraph.add_run(item["note"])
        document.add_heading("Presenter Notes", level=2)
        document.add_paragraph(item["speaker_notes"])
    document.save(path)


def main():
    root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--notebook", type=Path, default=root / "outputs" / "finance_demo.executed.ipynb")
    args = parser.parse_args()
    destination = root / "docs" / "presentation"
    content = json.loads((destination / "content.json").read_text(encoding="utf-8"))
    if not args.notebook.exists():
        parser.error("Run python scripts/execute_notebook.py first to produce the demonstration charts")
    figures = extract_figures(args.notebook, destination / "assets")
    (destination / "PROJECT_PRESENTATION.md").write_text(markdown_text(content), encoding="utf-8")
    build_word(content, figures, destination / "PROJECT_PRESENTATION.docx")
    build_powerpoint(content, figures, destination / "PROJECT_PRESENTATION.pptx")
    print(f"Generated {len(content['slides'])} slides in Markdown, DOCX, and PPTX: {destination}")


if __name__ == "__main__":
    main()